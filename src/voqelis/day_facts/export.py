"""Readable DOCX, A4 landscape PDF and phone-first A5 PDF exports."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

from .store import FactsEntry

TITLE = "ФАКТЫ ДНЯ И ОШИБОЧНЫЕ СУЖДЕНИЯ"
COLS = ("События", "Факты дня", "Ошибочные суждения")


def _numbered(items: tuple[str, ...]) -> str:
    return "\n".join(f"{index}. {item}" for index, item in enumerate(items, 1))


def _prepare(entries: list[FactsEntry], day: date, output: Path) -> None:
    if not entries:
        raise ValueError("No entries to export for the requested date")
    if any(entry.day != day for entry in entries):
        raise ValueError("Mixed dates in export")
    if output.suffix.lower() not in {".pdf", ".docx"}:
        raise ValueError("Unsupported export extension")
    output.parent.mkdir(parents=True, exist_ok=True)


def export_docx(entries: list[FactsEntry], day: date, output: Path) -> Path:
    """A4 landscape Word table. Styles and widths are stable in Word/LibreOffice."""
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    _prepare(entries, day, output)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.55)

    style = doc.styles["Normal"]
    style.font.name = "DejaVu Sans"
    style.font.size = Pt(9)
    style.font.color.rgb = RGBColor(37, 50, 67)
    style.paragraph_format.space_after = Pt(5)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(8)
    run = heading.add_run(TITLE)
    run.bold = True
    run.font.name = "DejaVu Sans"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor(28, 56, 81)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(19)
    run = sub.add_run(day.strftime("%d.%m.%Y"))
    run.font.size = Pt(11)

    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = (Cm(8.4), Cm(9.2), Cm(9.2))
    for i, label in enumerate(COLS):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        cell.text = label
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(28, 56, 81)
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "E9F0F4")
        tc_pr.append(shd)
    # Repeat table header on every Word page.
    header_row = table.rows[0]
    tr_pr = header_row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)

    for entry in entries:
        cells = table.add_row().cells
        texts = (entry.event, _numbered(entry.facts), _numbered(entry.judgments))
        for i, value in enumerate(texts):
            cells[i].width = widths[i]
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cells[i].text = value
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.line_spacing = 1.2
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Факты дня · личные записи").font.size = Pt(8)
    doc.save(output)
    return output



def export_mobile_docx(entries: list[FactsEntry], day: date, output: Path) -> Path:
    """Portrait A5 editable Word with one column for comfortable phone reading."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    _prepare(entries, day, output)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(14.8)
    section.page_height = Cm(21)
    section.left_margin = Cm(1.25)
    section.right_margin = Cm(1.25)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(37, 50, 67)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.2

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(28, 56, 81)

    date_line = doc.add_paragraph(day.strftime("%d.%m.%Y"))
    date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_line.paragraph_format.space_after = Pt(14)

    def heading(label: str, *, before: int = 9) -> None:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(label)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(28, 56, 81)

    for index, entry in enumerate(entries, 1):
        heading(f"Ситуация {index}", before=14 if index > 1 else 2)
        heading("Событие")
        doc.add_paragraph(entry.event)
        heading("Факты дня")
        for number, fact in enumerate(entry.facts, 1):
            doc.add_paragraph(f"{number}. {fact}")
        heading("Ошибочные суждения")
        for number, judgment in enumerate(entry.judgments, 1):
            doc.add_paragraph(f"{number}. {judgment}")

    doc.save(output)
    return output


def _pdf_styles():
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    name = "DayFactsDejaVu"
    bold_name = "DayFactsDejaVuBold"
    font_root = Path("/usr/share/fonts/truetype/dejavu")
    regular = font_root / "DejaVuSans.ttf"
    bold = font_root / "DejaVuSans-Bold.ttf"
    if not regular.is_file() or not bold.is_file():
        raise FileNotFoundError("DejaVu fonts with Cyrillic are required")
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(name, str(regular)))
    if bold_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(bold_name, str(bold)))

    ink = colors.HexColor("#26374A")
    title = ParagraphStyle(
        "dayfacts_title", fontName=bold_name, fontSize=15.5,
        leading=21, textColor=colors.HexColor("#1C3851"),
        alignment=TA_CENTER, spaceAfter=11,
    )
    date_style = ParagraphStyle(
        "dayfacts_date", fontName=name, fontSize=10, leading=14,
        textColor=ink, alignment=TA_CENTER, spaceAfter=20,
    )
    body = ParagraphStyle(
        "dayfacts_body", fontName=name, fontSize=8.5, leading=13,
        textColor=ink, alignment=TA_LEFT,
    )
    th = ParagraphStyle(
        "dayfacts_th", parent=body, fontName=bold_name,
        fontSize=9, leading=13,
    )
    mobile = ParagraphStyle(
        "dayfacts_mobile", fontName=name, fontSize=10,
        leading=16, textColor=ink, spaceAfter=7, wordWrap="CJK",
    )
    mobile_title = ParagraphStyle(
        "dayfacts_mobile_title", parent=title, fontSize=15, leading=22,
    )
    label = ParagraphStyle(
        "dayfacts_label", fontName=bold_name, fontSize=10,
        leading=15, textColor=colors.HexColor("#1C3851"),
        spaceBefore=8, spaceAfter=5,
    )
    return title, date_style, body, th, mobile, mobile_title, label


def _paragraph(text: str, style):
    from reportlab.platypus import Paragraph

    return Paragraph(escape(text).replace("\n", "<br/>"), style)


def _footer(canvas, document) -> None:
    from reportlab.lib import colors

    canvas.saveState()
    canvas.setFont("DayFactsDejaVu", 7.5)
    canvas.setFillColor(colors.HexColor("#62758A"))
    canvas.drawCentredString(
        document.pagesize[0] / 2, 21,
        f"Факты дня · страница {document.page}",
    )
    canvas.restoreState()


def export_pdf(
    entries: list[FactsEntry], day: date, output: Path, *, mobile: bool = False
) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, A5, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        HRFlowable,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    _prepare(entries, day, output)
    title, date_style, body, th, mobile_body, mobile_title, label = _pdf_styles()
    pagesize = A5 if mobile else landscape(A4)
    margins = (13 * mm) if mobile else (16 * mm)
    doc = SimpleDocTemplate(
        str(output), pagesize=pagesize,
        leftMargin=margins, rightMargin=margins,
        topMargin=17 * mm, bottomMargin=18 * mm,
        title=TITLE, author="Voqelis",
        pageCompression=1,
    )
    story = [
        _paragraph(TITLE, mobile_title if mobile else title),
        _paragraph(day.strftime("%d.%m.%Y"), date_style),
    ]
    if mobile:
        # A5 portrait: reading a wide 3-column table on a phone is not practical.
        # Each situation retains all three fields, in their original order.
        for position, entry in enumerate(entries, 1):
            # Avoid KeepTogether on unbounded dictated paragraphs:
            # an individual event may span several phone pages.
            story.append(_paragraph(f"Ситуация {position}", label))
            story.append(_paragraph(f"Событие: {entry.event}", mobile_body))
            story.append(_paragraph("Факты дня", label))
            for index, fact in enumerate(entry.facts, 1):
                story.append(_paragraph(f"{index}. {fact}", mobile_body))
            story.append(_paragraph("Ошибочные суждения", label))
            for index, judgment in enumerate(entry.judgments, 1):
                story.append(_paragraph(f"{index}. {judgment}", mobile_body))
            story.append(Spacer(1, 7 * mm))
            if position < len(entries):
                story.append(HRFlowable(
                    width="100%", thickness=0.65,
                    color=colors.HexColor("#D5DEE8"),
                ))
    else:
        available = pagesize[0] - 2 * margins
        col_widths = [available * ratio for ratio in (0.31, 0.345, 0.345)]
        rows = [[_paragraph(x, th) for x in COLS]]
        for entry in entries:
            rows.append([
                _paragraph(entry.event, body),
                _paragraph(_numbered(entry.facts), body),
                _paragraph(_numbered(entry.judgments), body),
            ])
        table = Table(
            rows, colWidths=col_widths, repeatRows=1,
            splitByRow=1, splitInRow=1, hAlign="LEFT",
        )
        table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9F0F4")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F8FAFC")]),
            ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#CAD6E1")),
            ("LEFTPADDING", (0, 0), (-1, -1), 9),
            ("RIGHTPADDING", (0, 0), (-1, -1), 9),
            ("TOPPADDING", (0, 0), (-1, -1), 10),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ]))
        story.append(table)
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return output
