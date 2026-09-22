from __future__ import annotations

from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

from .config import PlannerConfig
from .models import PlanItem
from .render import _rows_for_plan
from .scheduler import fmt_time
from .store import PlannerStore


def _actual_text(review) -> str:
    if not review:
        return ""
    parts = [review.activity] if review.activity else []
    if review.feelings:
        parts.append(", ".join(review.feelings))
    return "\n".join(parts)


def build_docx(user_id: int, day: date, store: PlannerStore, config: PlannerConfig, output: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    rows = _rows_for_plan(store.plan_items(user_id, day), config)
    reviews = {x.plan_item.id: x for x in store.reviews(user_id, day)}

    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"План дня — {day.strftime('%d.%m.%Y')}")
    run.bold = True
    run.font.size = Pt(14)

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for index, header in enumerate(["", "Время", "Запланированные дела", "Зачем мне это нужно", "Дела, которыми я занимался / чувства"]):
        table.rows[0].cells[index].text = header

    meta: list[tuple[PlanItem | None, int]] = []
    for start, end, item in rows:
        cells = table.add_row().cells
        review = reviews.get(item.id) if item else None
        cells[0].text = review.status if review else ""
        cells[1].text = f"{fmt_time(start)}–{fmt_time(end)}"
        cells[2].text = item.title if item else ""
        cells[3].text = item.why if item else ""
        cells[4].text = _actual_text(review)
        meta.append((item, len(table.rows) - 1))

    index = 0
    while index < len(meta):
        item, first = meta[index]
        if item is None:
            index += 1
            continue
        end = index + 1
        while end < len(meta) and meta[end][0] is item:
            end += 1
        if end - index > 1:
            last = meta[end - 1][1]
            for column in (0, 2, 3, 4):
                table.cell(first, column).merge(table.cell(last, column))
        index = end

    day_review = store.day_review(user_id, day)
    if day_review:
        doc.add_paragraph("Что бы я изменил, если бы следовал рекомендации по оздоровлению?")
        doc.add_paragraph(day_review.what_would_change or "")
        doc.add_paragraph("Признаки срыва:")
        doc.add_paragraph(", ".join(day_review.relapse_signs))

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def build_pdf(user_id: int, day: date, store: PlannerStore, config: PlannerConfig, output: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    font_name = "Helvetica"
    if Path(font_path).exists():
        pdfmetrics.registerFont(TTFont("DejaVu", font_path))
        font_name = "DejaVu"

    rows = _rows_for_plan(store.plan_items(user_id, day), config)
    reviews = {x.plan_item.id: x for x in store.reviews(user_id, day)}
    body = ParagraphStyle(
        "planner_body", parent=getSampleStyleSheet()["BodyText"],
        fontName=font_name, fontSize=7, leading=9,
    )
    title = ParagraphStyle(
        "planner_title", parent=getSampleStyleSheet()["Title"],
        fontName=font_name, fontSize=14, alignment=TA_CENTER,
    )

    data = [[Paragraph(x, body) for x in ["", "Время", "Запланированные дела", "Зачем мне это нужно", "Дела, которыми я занимался / чувства"]]]
    row_items: list[PlanItem | None] = []
    for start, end, item in rows:
        review = reviews.get(item.id) if item else None
        data.append([
            Paragraph(review.status if review else "", body),
            Paragraph(f"{fmt_time(start)}–{fmt_time(end)}", body),
            Paragraph(item.title if item else "", body),
            Paragraph(item.why if item else "", body),
            Paragraph(escape(_actual_text(review)).replace("\n", "<br/>"), body),
        ])
        row_items.append(item)

    spans = []
    index = 0
    while index < len(row_items):
        item = row_items[index]
        if item is None:
            index += 1
            continue
        end = index + 1
        while end < len(row_items) and row_items[end] is item:
            end += 1
        if end - index > 1:
            for column in (0, 2, 3, 4):
                spans.append(("SPAN", (column, index + 1), (column, end)))
        index = end

    doc = SimpleDocTemplate(
        str(output), pagesize=landscape(A4),
        leftMargin=18, rightMargin=18, topMargin=24, bottomMargin=24,
    )
    table = Table(data, colWidths=[28, 58, 170, 150, 300], repeatRows=1)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        *spans,
    ]))

    story = [Paragraph(f"План дня — {day.strftime('%d.%m.%Y')}", title), table]
    day_review = store.day_review(user_id, day)
    if day_review:
        story.extend([
            Paragraph("Что бы я изменил, если бы следовал рекомендации по оздоровлению?", body),
            Paragraph(escape(day_review.what_would_change or ""), body),
            Paragraph("Признаки срыва:", body),
            Paragraph(escape(", ".join(day_review.relapse_signs)), body),
        ])
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.build(story)
    return output
