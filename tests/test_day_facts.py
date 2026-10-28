"""Day Facts: state-machine integrity, boundary correction, exports and isolation."""

from __future__ import annotations

import sqlite3
import stat
from datetime import date
from pathlib import Path

import pytest

from voqelis.day_facts.export import TITLE, export_docx, export_mobile_docx, export_pdf
from voqelis.day_facts.sentences import (
    split_explicit,
    split_sentences,
    split_voice_items,
)
from voqelis.day_facts.service import DayFactsService
from voqelis.day_facts.store import DayFactsStore


@pytest.fixture
def store(tmp_path: Path):
    db = DayFactsStore(tmp_path / "day_facts.sqlite3")
    yield db
    db.close()


def fill_record(store: DayFactsStore, user: int, day: date):
    draft = store.begin(user, day)
    assert draft.stage == "event"
    draft = store.accept(
        user, draft.token, "Гуляли с анонимным другом, он был без настроения."
    )
    assert draft is not None and draft.stage == "facts"
    draft = store.accept(
        user, draft.token,
        "Друг общался раздражённо, был в своих мыслях. Мы разошлись по домам.",
        raw_text="друг общался раздраженно был в своих мыслях мы разошлись по домам",
    )
    assert draft is not None
    assert len(draft.facts) == 2
    draft = store.next(user, draft.token)
    assert draft is not None and draft.stage == "judgments"
    draft = store.accept(
        user, draft.token,
        "Я как-то не так общаюсь. Из меня друг так себе."
    )
    assert draft is not None and len(draft.judgments) == 2
    draft = store.next(user, draft.token)
    assert draft is not None and draft.stage == "review"
    return draft


def test_sentence_split_uses_full_stops_not_commas() -> None:
    assert split_sentences(
        "Друг общался раздражённо, был в своих мыслях. Мы разошлись по домам."
    ) == (
        "Друг общался раздражённо, был в своих мыслях.",
        "Мы разошлись по домам.",
    )
    assert split_sentences("Я не так общаюсь? Я плохой друг!") == (
        "Я не так общаюсь?", "Я плохой друг!"
    )
    assert split_sentences("факт без точки другой факт без точки") == (
        "факт без точки другой факт без точки",
    )
    assert split_sentences(
        "Друг смотрел в окно следующий факт мы разошлись по домам"
    ) == ("Друг смотрел в окно", "мы разошлись по домам")
    assert split_sentences(
        "Я испортил встречу следующее суждение я плохой друг"
    ) == ("Я испортил встречу", "я плохой друг")



# Voice-specific capture: punctuation is presentation, not an implicit
# instruction to split the user's personal diary into multiple records.
def test_voice_split_preserves_internal_periods_and_respects_spoken_markers() -> None:
    text = "Друг говорил тихо. Потом посмотрел в окно."
    assert split_voice_items(text) == (text,)
    assert split_voice_items(
        "Друг говорил тихо. Следующий факт мы разошлись по домам."
    ) == ("Друг говорил тихо.", "мы разошлись по домам.")
    assert split_voice_items(
        "Я расстроился? Следующее суждение я плохо поговорил."
    ) == ("Я расстроился?", "я плохо поговорил.")
    with pytest.raises(ValueError):
        split_voice_items("следующий факт")
    with pytest.raises(ValueError):
        split_voice_items("Событие. Следующий факт")


def test_voice_markers_are_section_specific_and_resist_stray_punctuation(store) -> None:
    assert split_voice_items(
        "Друг ушёл. Следующий. факт: Я пошёл домой.", field="facts"
    ) == ("Друг ушёл.", "Я пошёл домой.")
    assert split_voice_items(
        "Следующее суждение меня напугало.", field="facts"
    ) == ("Следующее суждение меня напугало.",)
    assert split_voice_items(
        "Следующий факт меня удивил.", field="judgments"
    ) == ("Следующий факт меня удивил.",)
    with pytest.raises(ValueError, match="Unknown Day Facts section"):
        split_voice_items("Любой текст", field="event")

    draft = store.begin(92, date(2026, 10, 26))
    draft = store.accept(92, draft.token, "Встреча")
    draft = store.accept(
        92, draft.token, "Следующее суждение меня удивило.",
        voice_dictation=True,
    )
    assert draft.facts == ("Следующее суждение меня удивило.",)
    draft = store.next(92, draft.token)
    draft = store.accept(
        92, draft.token, "Следующий факт меня расстроил.",
        voice_dictation=True,
    )
    assert draft.judgments == ("Следующий факт меня расстроил.",)


def test_voice_accept_adds_exactly_one_fact_and_one_judgment(store) -> None:
    draft = store.begin(88, date(2026, 10, 26))
    draft = store.accept(88, draft.token, "Встреча с другом")
    old = draft.token
    text = "Друг говорил тихо. Потом посмотрел в окно."
    draft = store.accept(
        88, old, text, raw_text="друг говорил тихо потом посмотрел в окно",
        voice_dictation=True,
    )
    assert draft is not None
    assert draft.facts == (text,)
    assert store.accept(88, old, "устаревший голос", voice_dictation=True) is None

    draft = store.accept(
        88, draft.token, "Мы попрощались.", voice_dictation=True,
    )
    assert draft.facts == (text, "Мы попрощались.")
    draft = store.next(88, draft.token)
    draft = store.accept(
        88, draft.token, "Мне показалось, что я плохо общался. "
        "Я испортил встречу.",
        voice_dictation=True,
    )
    assert draft.judgments == (
        "Мне показалось, что я плохо общался. Я испортил встречу.",
    )
    draft = store.next(88, draft.token)
    saved = store.save(88, draft.token)
    assert saved is not None
    assert saved.facts == (text, "Мы попрощались.")
    raw = store.db.execute(
        "SELECT raw_inputs_json FROM day_facts_entries WHERE id=?", (saved.id,)
    ).fetchone()[0]
    assert "друг говорил тихо потом посмотрел в окно" in raw


def test_explicit_spoken_marker_is_the_only_voice_split(store) -> None:
    draft = store.begin(89, date(2026, 10, 26))
    draft = store.accept(89, draft.token, "Встреча")
    draft = store.accept(
        89, draft.token,
        "Друг отвечал кратко. Следующий факт мы пошли домой. "
        "Следующий факт наступил вечер.",
        voice_dictation=True,
    )
    assert draft.facts == (
        "Друг отвечал кратко.",
        "мы пошли домой.",
        "наступил вечер.",
    )
    draft = store.next(89, draft.token)
    draft = store.accept(
        89, draft.token,
        "Мне было грустно. Следующее суждение я плохой собеседник.",
        voice_dictation=True,
    )
    assert draft.judgments == ("Мне было грустно.", "я плохой собеседник.")


def test_malformed_voice_marker_cannot_modify_draft(store) -> None:
    draft = store.begin(90, date(2026, 10, 26))
    draft = store.accept(90, draft.token, "Событие")
    with pytest.raises(ValueError, match="spoken item"):
        store.accept(
            90, draft.token, "История. Следующий факт", voice_dictation=True,
        )
    assert store.draft(90) == draft


def test_text_input_preserves_existing_period_based_splitting(store) -> None:
    draft = store.begin(91, date(2026, 10, 26))
    draft = store.accept(91, draft.token, "Событие")
    draft = store.accept(91, draft.token, "Один факт. Второй факт.")
    assert draft.facts == ("Один факт.", "Второй факт.")


def test_explicit_split_allows_boundary_correction() -> None:
    assert split_explicit(
        "Друг был раздражён | Мы разошлись по домам"
    ) == ("Друг был раздражён", "Мы разошлись по домам")
    with pytest.raises(ValueError):
        split_explicit("только один")
    with pytest.raises(ValueError):
        split_explicit("два |  ")


def test_complete_example_saves_one_event_two_facts_two_judgments(store):
    day = date(2026, 10, 26)
    draft = fill_record(store, 7, day)
    note = store.save(7, draft.token)
    assert note is not None
    assert note.event == "Гуляли с анонимным другом, он был без настроения."
    assert len(note.facts) == 2
    assert note.facts[0] == "Друг общался раздражённо, был в своих мыслях."
    assert note.facts[1] == "Мы разошлись по домам."
    assert note.judgments == (
        "Я как-то не так общаюсь.", "Из меня друг так себе."
    )
    assert store.draft(7) is None
    assert [item.id for item in store.for_day(7, day)] == [note.id]
    raw_row = store.db.execute(
        "SELECT raw_inputs_json FROM day_facts_entries WHERE id=?", (note.id,)
    ).fetchone()
    assert "друг общался раздраженно" in raw_row[0]


def test_unfinished_or_empty_draft_is_never_saved(store):
    draft = store.begin(1, date(2026, 10, 26))
    assert store.save(1, draft.token) is None
    with pytest.raises(ValueError):
        store.next(1, draft.token)
    assert store.for_day(1, date(2026, 10, 26)) == []


def test_second_begin_preserves_pending_draft(store):
    day = date(2026, 10, 26)
    a = store.begin(1, day)
    b = store.begin(1, date(2026, 10, 27))
    assert b == a


def test_replay_stale_voice_is_rejected_and_does_not_overwrite(store):
    draft = store.begin(1, date(2026, 10, 26))
    old_token = draft.token
    updated = store.accept(1, draft.token, "Встреча с другом")
    assert updated
    assert store.accept(1, old_token, "чужое сообщение") is None
    assert store.draft(1) == updated
    assert store.for_day(1, date(2026, 10, 26)) == []


def test_new_draft_invalidates_old_token_after_cancel(store):
    day = date(2026, 10, 26)
    old = store.begin(1, day)
    assert store.cancel(1)
    fresh = store.begin(1, day)
    assert fresh.token != old.token
    assert store.accept(1, old.token, "не сохранять") is None


def test_next_requires_at_least_one_fact_and_judgment(store):
    d = store.begin(1, date(2026, 10, 26))
    d = store.accept(1, d.token, "Событие")
    with pytest.raises(ValueError):
        store.next(1, d.token)
    d = store.accept(1, d.token, "Факт.")
    d = store.next(1, d.token)
    with pytest.raises(ValueError):
        store.next(1, d.token)


def test_manual_split_and_correction_even_at_review_stage(store):
    d = fill_record(store, 1, date(2026, 10, 26))
    before = d.token
    d = store.correct(
        1, d.token, "judgments", 1,
        "Я неправильно общаюсь | Я всё испортил",
        split=True,
    )
    assert d
    assert len(d.judgments) == 3
    assert d.judgments[:2] == ("Я неправильно общаюсь", "Я всё испортил")
    assert store.correct(1, before, "facts", 1, "подмена") is None
    d = store.correct(1, d.token, "event", 1, "Другое событие.")
    assert d.event == "Другое событие."
    with pytest.raises(ValueError):
        store.correct(1, d.token, "facts", 99, "нет")
    with pytest.raises(ValueError):
        store.correct(1, d.token, "event", 1, "a | b", split=True)


def test_user_isolation_dates_and_history(store):
    day = date(2026, 10, 26)
    d = fill_record(store, 1, day)
    saved = store.save(1, d.token)
    assert saved
    assert store.get(2, saved.id) is None
    assert store.for_day(2, day) == []
    assert store.history(2) == []
    assert store.history(1) == [(day.isoformat(), 1)]


def test_independent_sqlite_file_does_not_modify_planner(tmp_path):
    planner = tmp_path / "planner.sqlite3"
    with sqlite3.connect(planner) as conn:
        conn.execute("CREATE TABLE preserved(x INTEGER)")
        conn.execute("INSERT INTO preserved VALUES(15)")
    ours = DayFactsStore(tmp_path / "day_facts.sqlite3")
    ours.begin(1, date(2026, 10, 26))
    ours.close()
    with sqlite3.connect(f"file:{planner}?mode=ro", uri=True) as db:
        assert db.execute("SELECT x FROM preserved").fetchone() == (15,)
        names = {row[0] for row in db.execute(
            "SELECT name FROM sqlite_master WHERE type='table'"
        )}
        assert "day_facts_drafts" not in names


def test_persistence_after_restart(tmp_path):
    path = tmp_path / "day_facts.sqlite3"
    one = DayFactsStore(path)
    review = fill_record(one, 1, date(2026, 10, 26))
    saved = one.save(1, review.token)
    one.close()
    two = DayFactsStore(path)
    try:
        assert two.get(1, saved.id) == saved
        assert two.db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    finally:
        two.close()


def test_service_review_displays_individual_items(store):
    service = DayFactsService(store)
    draft = fill_record(store, 1, date(2026, 10, 26))
    rendered = service.render(draft)
    assert "1. Друг общался" in rendered
    assert "2. Мы разошлись" in rendered
    assert "1. Я как-то" in rendered
    assert "2. Из меня" in rendered
    store.save(1, draft.token)
    assert "Ситуация 1" in service.day_message(1, date(2026, 10, 26))
    assert "26.10.2026" in service.history_message(1)


def test_docx_pdf_and_phone_pdf_include_cyrillic_and_three_fields(store, tmp_path):
    d = fill_record(store, 1, date(2026, 10, 26))
    store.save(1, d.token)
    entries = store.for_day(1, date(2026, 10, 26))
    word = export_docx(entries, d.day, tmp_path / "sample.docx")
    pdf = export_pdf(entries, d.day, tmp_path / "sample.pdf")
    phone = export_pdf(entries, d.day, tmp_path / "phone.pdf", mobile=True)

    from docx import Document

    opened = Document(word)
    assert TITLE in "\n".join(p.text for p in opened.paragraphs)
    assert len(opened.tables) == 1
    assert len(opened.tables[0].columns) == 3
    assert len(opened.tables[0].rows) == 2
    assert "Друг общался" in opened.tables[0].cell(1, 1).text
    assert "2. Из меня" in opened.tables[0].cell(1, 2).text
    assert pdf.read_bytes().startswith(b"%PDF")
    assert phone.read_bytes().startswith(b"%PDF")
    assert all(p.stat().st_size > 2000 for p in (word, pdf, phone))



def test_phone_word_docx_is_portrait_editable_and_keeps_all_fields(store, tmp_path):
    day = date(2026, 10, 26)
    draft = fill_record(store, 1, day)
    store.save(1, draft.token)
    entries = store.for_day(1, day)

    mobile = export_mobile_docx(entries, day, tmp_path / "phone.docx")

    from docx import Document

    doc = Document(mobile)
    assert not doc.tables
    section = doc.sections[0]
    assert section.page_width < section.page_height
    assert abs(section.page_width.cm - 14.8) < 0.1
    assert doc.styles["Normal"].font.size.pt == 11

    text = "\n".join(p.text for p in doc.paragraphs)
    assert TITLE in text
    assert "26.10.2026" in text
    assert "Ситуация 1" in text
    assert "Событие" in text
    assert "Факты дня" in text
    assert "Ошибочные суждения" in text
    assert entries[0].event in text
    assert all(f"{i}. {fact}" in text for i, fact in enumerate(entries[0].facts, 1))
    assert all(
        f"{i}. {judgment}" in text
        for i, judgment in enumerate(entries[0].judgments, 1)
    )
    assert mobile.stat().st_size > 2000

    with pytest.raises(ValueError):
        export_mobile_docx(entries, date(2026, 10, 27), tmp_path / "wrong.docx")


def test_phone_word_docx_handles_multiple_situations(store, tmp_path):
    day = date(2026, 10, 26)
    for number in range(2):
        draft = store.begin(1, day)
        draft = store.accept(1, draft.token, f"Событие номер {number + 1}")
        draft = store.accept(1, draft.token, f"Факт номер {number + 1}")
        draft = store.next(1, draft.token)
        draft = store.accept(1, draft.token, f"Суждение номер {number + 1}")
        draft = store.next(1, draft.token)
        store.save(1, draft.token)

    from docx import Document

    output = export_mobile_docx(store.for_day(1, day), day, tmp_path / "multi.docx")
    doc = Document(output)
    content = "\n".join(p.text for p in doc.paragraphs)
    assert "Ситуация 1" in content and "Ситуация 2" in content
    assert "Событие номер 1" in content and "Событие номер 2" in content
    assert not doc.tables


def test_many_day_entries_export_more_than_one_page(store, tmp_path):
    day = date(2026, 10, 26)
    for i in range(25):
        d = store.begin(1, day)
        d = store.accept(1, d.token, f"Ситуация номер {i}")
        d = store.accept(1, d.token, "Проверяем факт номер один. Факт номер два.")
        d = store.next(1, d.token)
        d = store.accept(1, d.token, "Это первая мысль. Это вторая мысль.")
        d = store.next(1, d.token)
        store.save(1, d.token)
    entries = store.for_day(1, day)
    assert len(entries) == 25
    output = export_pdf(entries, day, tmp_path / "many.pdf")
    output_phone = export_pdf(entries, day, tmp_path / "many-mobile.pdf", mobile=True)
    assert output.stat().st_size > 3000
    assert output_phone.stat().st_size > 3000


def test_invalid_export_inputs_fail_closed(store, tmp_path):
    with pytest.raises(ValueError):
        export_pdf([], date(2026, 10, 26), tmp_path / "empty.pdf")
    d = fill_record(store, 1, date(2026, 10, 26))
    item = store.save(1, d.token)
    with pytest.raises(ValueError):
        export_docx([item], date(2026, 10, 27), tmp_path / "wrong.docx")


# Telegram/ASR regression tests: model text is presentation, not a fact source of truth.
import asyncio
from types import SimpleNamespace

from voqelis.bot import run_worker
from voqelis.day_facts.bot import create_day_facts_router, export_markup, facts_markup
from voqelis.domain import AudioJob, TranscriptionResult


class _FakeBot:
    def __init__(self):
        self.replies = []

    async def send_chat_action(self, *_args):
        pass

    async def send_message(self, _chat_id, body, **kwargs):
        self.replies.append((body, kwargs))


class _FakeQueue:
    def __init__(self, job):
        self.job = job
        self.n = 0

    async def get(self):
        self.n += 1
        if self.n > 1:
            raise asyncio.CancelledError
        return self.job

    async def release(self, _id):
        pass

    def task_done(self):
        pass


class _FakeASR:
    async def transcribe(self, _path):
        return TranscriptionResult(
            text="друг общался раздраженно был в своих мыслях мы разошлись по домам",
            pause_aware_text=None,
            language="ru",
            language_probability=0,
            duration_seconds=3,
            duration_after_vad_seconds=3,
            processing_seconds=0.1,
        )


class _FakePunctuation:
    async def format(self, source, *, fallback_text):
        assert source.startswith("друг общался")
        return (
            "Друг общался раздражённо, был в своих мыслях. "
            "Мы разошлись по домам."
        )


async def _worker_case(tmp_path, monkeypatch, service, token):
    from voqelis import bot as bot_module

    path = tmp_path / "sample.ogg"
    path.write_bytes(b"x")
    monkeypatch.setattr(bot_module, "probe_duration_seconds", lambda _path: 3)
    bot = _FakeBot()
    job = AudioJob(
        chat_id=7, reply_to_message_id=10, user_id=1,
        file_path=path, day_facts_token=token,
    )
    with pytest.raises(asyncio.CancelledError):
        await run_worker(
            bot=bot,
            queue=_FakeQueue(job),
            transcriber=_FakeASR(),
            settings=SimpleNamespace(max_audio_seconds=90),
            day_facts=service,
            punctuation_processor=_FakePunctuation(),
        )
    return bot.replies


@pytest.mark.asyncio
async def test_voice_to_facts_preserves_original_words_as_single_item(
    store, tmp_path, monkeypatch
):
    service = DayFactsService(store)
    draft = store.begin(1, date(2026, 10, 26))
    draft = store.accept(1, draft.token, "Прогулка с другом.")
    result = await _worker_case(tmp_path, monkeypatch, service, draft.token)
    after = store.draft(1)
    assert after is not None and after.stage == "facts"
    assert after.facts == (
        ("Друг общался раздражённо, был в своих мыслях. "
         "Мы разошлись по домам."),
    )
    assert "Друг общался" in result[0][0]
    assert "1. Друг общался" in result[0][0]
    assert "2. Мы разошлись" not in result[0][0]
    assert result[0][1]["reply_markup"] is not None
    raw = store.db.execute(
        "SELECT raw_inputs_json FROM day_facts_drafts WHERE user_id=1"
    ).fetchone()[0]
    assert "друг общался раздраженно" in raw


@pytest.mark.asyncio
async def test_stale_voice_queued_before_step_change_cannot_mutate_facts(
    store, tmp_path, monkeypatch
):
    service = DayFactsService(store)
    draft = store.begin(1, date(2026, 10, 26))
    stale = draft.token
    changed = store.accept(1, stale, "Уже другое событие.")
    assert changed is not None
    result = await _worker_case(tmp_path, monkeypatch, service, stale)
    assert "не сохранено" in result[0][0]
    assert store.draft(1) == changed


@pytest.mark.asyncio
async def test_invalid_spoken_marker_worker_feedback_is_not_overlong(
    store, tmp_path, monkeypatch
):
    service = DayFactsService(store)
    draft = store.begin(1, date(2026, 10, 26))
    draft = store.accept(1, draft.token, "Событие.")

    def reject_empty_item(*_args, **_kwargs):
        raise ValueError("Each spoken item must contain text")

    monkeypatch.setattr(service, "accept", reject_empty_item)
    replies = await _worker_case(tmp_path, monkeypatch, service, draft.token)
    assert "разделители" in replies[0][0]
    assert "слишком длинный" not in replies[0][0]
    assert store.draft(1) == draft


def test_day_facts_router_has_command_and_keyboard_handlers(store, tmp_path):
    service = DayFactsService(store)
    planner = SimpleNamespace(store=SimpleNamespace(session=lambda _id: None))
    notes = SimpleNamespace(pending_voice=lambda _id: False)
    router = create_day_facts_router(
        service=service, allowed_user_ids=frozenset({1}), planner=planner,
        notes=notes, export_dir=tmp_path,
    )
    callbacks = {h.callback.__name__ for h in router.message.handlers}
    assert {
        "on_start", "on_cancel", "on_next", "on_edit", "on_split",
        "on_history", "on_day", "on_export", "on_field_text",
    } <= callbacks
    assert len(router.callback_query.handlers) == 1
    d = store.begin(1, date(2026, 10, 26))
    assert len(facts_markup(d).inline_keyboard) == 1
    export_buttons = [
        button.callback_data
        for row in export_markup(date(2026, 10, 26)).inline_keyboard
        for button in row
    ]
    assert export_buttons == [
        "df:export:word:2026-10-26",
        "df:export:pdf:2026-10-26",
        "df:export:word_mobile:2026-10-26",
        "df:export:mobile:2026-10-26",
    ]


@pytest.mark.parametrize("kind", ("word", "word_mobile"))
@pytest.mark.asyncio
async def test_export_command_executes_real_threaded_docx_without_sqlite_error(
    store, tmp_path, monkeypatch, kind
):
    service = DayFactsService(store)
    monkeypatch.setattr(service, "today", lambda: date(2026, 10, 26))
    draft = fill_record(store, 1, date(2026, 10, 26))
    store.save(1, draft.token)

    class FakeMessage:
        text = f"/facts_export {kind}"
        from_user = SimpleNamespace(id=1)
        chat = SimpleNamespace(type="private")

        def __init__(self):
            self.documents = []
            self.messages = []

        async def answer(self, value, **kwargs):
            self.messages.append(value)

        async def answer_document(self, attachment, **kwargs):
            # FSInputFile is lazily opened; inspect while still inside the
            # export handler, before the generated temporary file is removed.
            self.documents.append((attachment.path, Path(attachment.path).read_bytes()))

    planner = SimpleNamespace(store=SimpleNamespace(session=lambda _id: None))
    notes = SimpleNamespace(pending_voice=lambda _id: False)
    router = create_day_facts_router(
        service=service, allowed_user_ids=frozenset({1}), planner=planner,
        notes=notes, export_dir=tmp_path,
    )
    handler = next(
        h.callback for h in router.message.handlers
        if h.callback.__name__ == "on_export"
    )
    msg = FakeMessage()
    await handler(msg)
    assert len(msg.documents) == 1
    assert msg.documents[0][1][:2] == b"PK"
    assert not Path(msg.documents[0][0]).exists()
    assert msg.messages == []



def test_very_long_event_can_export_across_pdf_pages(store, tmp_path):
    day = date(2026, 10, 26)
    d = store.begin(1, day)
    # The event is one text field; it can be longer than a single PDF page.
    lengthy = ("Говорили о погоде и планах на неделю. " * 105).strip()
    d = store.accept(1, d.token, lengthy)
    d = store.accept(1, d.token, "Факт один. Факт два.")
    d = store.next(1, d.token)
    d = store.accept(1, d.token, "Первая мысль. Вторая мысль.")
    d = store.next(1, d.token)
    assert store.save(1, d.token)
    entries = store.for_day(1, day)
    assert export_pdf(entries, day, tmp_path / "long.pdf").stat().st_size > 3000
    assert export_pdf(
        entries, day, tmp_path / "long-mobile.pdf", mobile=True,
    ).stat().st_size > 3000



@pytest.mark.asyncio
async def test_stale_day_facts_voice_is_not_mistaken_for_planner_task(
    store, tmp_path, monkeypatch
):
    from voqelis import bot as bot_module

    service = DayFactsService(store)
    draft = store.begin(1, date(2026, 10, 26))
    path = tmp_path / "voice-plan-race.ogg"
    path.write_bytes(b"sample")
    monkeypatch.setattr(bot_module, "probe_duration_seconds", lambda _: 3)

    class GuardPlanner:
        config = SimpleNamespace(timezone="Europe/Moscow")
        store = SimpleNamespace(session=lambda _id: {
            "state": "planning", "target_day": "2026-10-26",
        })

        async def handle_text(self, *args, **kwargs):
            raise AssertionError("Must never create a Planner task from a facts voice")

    bot = _FakeBot()
    job = AudioJob(
        chat_id=7, reply_to_message_id=10, user_id=1, file_path=path,
        day_facts_token=draft.token,
    )
    with pytest.raises(asyncio.CancelledError):
        await run_worker(
            bot=bot, queue=_FakeQueue(job), transcriber=_FakeASR(),
            settings=SimpleNamespace(max_audio_seconds=90), planner=GuardPlanner(),
            day_facts=service, punctuation_processor=_FakePunctuation(),
        )
    assert "не добавлено" in bot.replies[0][0]
    assert store.draft(1) == draft
    assert store.for_day(1, date(2026, 10, 26)) == []


@pytest.mark.asyncio
async def test_queued_note_keeps_note_intent_when_facts_started_later(
    store, tmp_path, monkeypatch
):
    from voqelis import bot as bot_module
    from voqelis.notes.service import NotesService
    from voqelis.notes.store import NotesStore

    notes_store = NotesStore(tmp_path / "notes.sqlite3")
    try:
        notes = NotesService(notes_store)
        notes.arm_voice(1)
        note_token = notes.claim_voice(1)
        assert note_token
        day_facts = DayFactsService(store)
        draft = store.begin(1, date(2026, 10, 26))

        path = tmp_path / "voice-notes-race.ogg"
        path.write_bytes(b"sample")
        monkeypatch.setattr(bot_module, "probe_duration_seconds", lambda _: 3)
        bot = _FakeBot()
        job = AudioJob(
            chat_id=7, reply_to_message_id=10, user_id=1, file_path=path,
            capture_token=note_token,
        )
        with pytest.raises(asyncio.CancelledError):
            await run_worker(
                bot=bot, queue=_FakeQueue(job), transcriber=_FakeASR(),
                settings=SimpleNamespace(max_audio_seconds=90),
                notes=notes, day_facts=day_facts,
                punctuation_processor=_FakePunctuation(),
            )
        assert len(notes_store.recent(1)) == 1
        assert "Заметка #1 сохранена" in bot.replies[0][0]
        assert store.draft(1) == draft
    finally:
        notes_store.close()


@pytest.mark.asyncio
async def test_plain_voice_remains_plain_when_facts_mode_started_after_admission(
    store, tmp_path, monkeypatch
):
    from voqelis import bot as bot_module

    service = DayFactsService(store)
    draft = store.begin(1, date(2026, 10, 26))
    path = tmp_path / "previous-voice.ogg"
    path.write_bytes(b"sample")
    monkeypatch.setattr(bot_module, "probe_duration_seconds", lambda _: 3)
    bot = _FakeBot()
    job = AudioJob(chat_id=7, reply_to_message_id=10, user_id=1, file_path=path)
    with pytest.raises(asyncio.CancelledError):
        await run_worker(
            bot=bot, queue=_FakeQueue(job), transcriber=_FakeASR(),
            settings=SimpleNamespace(max_audio_seconds=90),
            day_facts=service, punctuation_processor=_FakePunctuation(),
        )
    assert bot.replies[0][0].startswith("Друг общался")
    assert store.draft(1) == draft



def test_day_facts_db_permissions_and_symlink_rejection(tmp_path):
    private_db = tmp_path / "private.sqlite3"
    ours = DayFactsStore(private_db)
    try:
        assert stat.S_IMODE(private_db.stat().st_mode) == 0o600
    finally:
        ours.close()
    # Repair overly broad legacy permissions at next safe open.
    private_db.chmod(0o644)
    reopened = DayFactsStore(private_db)
    try:
        assert stat.S_IMODE(private_db.stat().st_mode) == 0o600
    finally:
        reopened.close()
    alias = tmp_path / "linked.sqlite3"
    alias.symlink_to(private_db)
    with pytest.raises(OSError):
        DayFactsStore(alias)


@pytest.mark.asyncio
async def test_export_document_is_never_world_readable(store, tmp_path, monkeypatch):
    service = DayFactsService(store)
    monkeypatch.setattr(service, "today", lambda: date(2026, 10, 26))
    draft = fill_record(store, 1, date(2026, 10, 26))
    store.save(1, draft.token)

    class FakeMessage:
        text = "/facts_export pdf"
        from_user = SimpleNamespace(id=1)
        chat = SimpleNamespace(type="private")

        def __init__(self):
            self.mode = None
            self.responses = []

        async def answer(self, value, **kwargs):
            self.responses.append(value)

        async def answer_document(self, attachment, **kwargs):
            self.mode = stat.S_IMODE(Path(attachment.path).stat().st_mode)
            assert Path(attachment.path).read_bytes().startswith(b"%PDF")

    planner = SimpleNamespace(store=SimpleNamespace(session=lambda _id: None))
    notes = SimpleNamespace(pending_voice=lambda _id: False)
    router = create_day_facts_router(
        service=service, allowed_user_ids=frozenset({1}), planner=planner,
        notes=notes, export_dir=tmp_path,
    )
    handler = next(
        h.callback for h in router.message.handlers
        if h.callback.__name__ == "on_export"
    )
    msg = FakeMessage()
    await handler(msg)
    assert msg.mode == 0o600
    assert not msg.responses



def test_home_keyboard_contains_modules_and_facts_has_own_menu() -> None:
    from voqelis.menu import facts_keyboard, main_keyboard
    from voqelis.planner.bot import planner_keyboard

    home = [item.text for row in main_keyboard(admin=True).keyboard for item in row]
    guest_home = [item.text for row in main_keyboard().keyboard for item in row]
    assert guest_home == ["📘 Факты дня", "✅ Список дел (скоро)"]
    facts = [item.text for row in facts_keyboard().keyboard for item in row]
    planner = [item.text for row in planner_keyboard().keyboard for item in row]

    assert "📘 Факты дня" in home
    assert "✅ Список дел (скоро)" in home
    assert "📝 Заметки" in home
    assert "📅 Планировщик" in home
    assert "🎙 Расшифровка" in home
    assert "📋 Текущие факты дня" not in home
    assert "🗓 История фактов" not in home
    assert "📄 Word" not in home
    assert "📱 Word для телефона" not in home
    assert "📋 Текущие факты дня" in facts
    assert "🗓 История фактов" in facts
    assert "⬅️ Главное меню" in facts
    assert "📅 Планирование дня" in planner
    assert "⬅️ Главное меню" in planner


def test_redo_facts_discards_duplicates_without_touching_event_or_judgments(store):
    day = date(2026, 10, 26)
    draft = fill_record(store, 17, day)
    before_facts = draft.facts
    old_token = draft.token
    restarted = store.reset_section(17, draft.token, "facts")
    assert restarted is not None
    assert restarted.stage == "facts"
    assert restarted.facts == ()
    assert restarted.judgments == draft.judgments
    assert restarted.event == draft.event
    assert restarted.token != old_token
    assert store.accept(17, old_token, "Старое голосовое") is None
    assert store.reset_section(17, old_token, "facts") is None
    redictated = store.accept(
        17, restarted.token,
        "Друг общался раздражённо, был в своих мыслях. Мы разошлись по домам.",
    )
    assert redictated is not None
    assert redictated.facts == before_facts
    assert len(redictated.facts) == 2
    to_judgments = store.next(17, redictated.token)
    assert to_judgments.judgments == draft.judgments
    back_to_review = store.next(17, to_judgments.token)
    saved = store.save(17, back_to_review.token)
    assert saved is not None
    assert saved.facts == before_facts
    assert saved.judgments == draft.judgments
    assert len(store.for_day(17, day)) == 1


def test_redo_judgments_does_not_erase_facts(store):
    draft = fill_record(store, 17, date(2026, 10, 26))
    updated = store.reset_section(17, draft.token, "judgments")
    assert updated.stage == "judgments"
    assert updated.facts == draft.facts
    assert updated.judgments == ()
    new = store.accept(17, updated.token, "Мне показалось, что я плохой друг.")
    assert new.judgments == ("Мне показалось, что я плохой друг.",)
    review = store.next(17, new.token)
    saved = store.save(17, review.token)
    assert saved.facts == draft.facts
    assert saved.judgments == ("Мне показалось, что я плохой друг.",)


def test_replace_whole_facts_list_fixes_comma_without_duplicate(store):
    draft = fill_record(store, 17, date(2026, 10, 26))
    fixed = store.replace_section(
        17, draft.token, "facts",
        "Друг общался раздражённо, был в своих мыслях. Мы разошлись по домам.",
    )
    assert fixed.stage == "review"
    assert fixed.facts == (
        "Друг общался раздражённо, был в своих мыслях.",
        "Мы разошлись по домам.",
    )
    assert fixed.judgments == draft.judgments
    assert fixed.token != draft.token
    assert store.replace_section(17, draft.token, "facts", "старое") is None
    saved = store.save(17, fixed.token)
    assert saved.facts == fixed.facts
    assert len(saved.facts) == 2


def test_reset_and_replace_scope_and_invalid_stage(store):
    a = store.begin(17, date(2026, 10, 26))
    with pytest.raises(ValueError):
        store.reset_section(17, a.token, "facts")
    with pytest.raises(ValueError):
        store.replace_section(17, a.token, "judgments", "текст")
    with pytest.raises(ValueError):
        store.reset_section(17, a.token, "event")
    assert store.reset_section(18, a.token, "facts") is None
    assert store.replace_section(18, a.token, "facts", "чужие слова") is None
    a = store.accept(17, a.token, "Событие")
    with pytest.raises(ValueError):
        store.reset_section(17, a.token, "judgments")
    with pytest.raises(ValueError):
        store.replace_section(17, a.token, "facts", " ")
    after = store.replace_section(
        17, a.token, "facts", "Первый факт. Второй факт."
    )
    assert after.facts == ("Первый факт.", "Второй факт.")
    assert after.stage == "facts"


@pytest.mark.asyncio
async def test_redo_callback_requires_current_token_and_redicts_cleanly(
    store, tmp_path
):
    service = DayFactsService(store)
    pending = fill_record(store, 17, date(2026, 10, 26))
    planner = SimpleNamespace(store=SimpleNamespace(session=lambda _id: None))
    notes = SimpleNamespace(pending_voice=lambda _id: False)
    router = create_day_facts_router(
        service=service, allowed_user_ids=frozenset({17}),
        planner=planner, notes=notes, export_dir=tmp_path,
    )
    on_callback = router.callback_query.handlers[0].callback

    class MessageStub:
        chat = SimpleNamespace(type="private")

        def __init__(self):
            self.responses = []

        async def answer(self, value, **kwargs):
            self.responses.append((value, kwargs))

    class CallbackStub:
        from_user = SimpleNamespace(id=17)

        def __init__(self, message, data):
            self.message = message
            self.data = data
            self.ack = []

        async def answer(self, value="", **kwargs):
            self.ack.append((value, kwargs))

    message = MessageStub()
    command = f"df:redo:facts:{pending.token}"
    assert command in [
        item.callback_data
        for line in facts_markup(pending).inline_keyboard for item in line
    ]
    callback = CallbackStub(message, command)
    await on_callback(callback)
    assert callback.ack
    after = store.draft(17)
    assert after.facts == ()
    assert after.judgments == pending.judgments
    assert after.stage == "facts"
    assert any("Факты очищены" in text for text, _ in message.responses)

    await on_callback(CallbackStub(message, command))
    assert store.draft(17) == after
    # An already queued ASR transcript can no longer re-add the old facts.
    assert store.accept(17, pending.token, "старое сообщение") is None


@pytest.mark.asyncio
async def test_day_facts_menu_routes_today_and_history_without_modifying_draft(
    store, tmp_path
):
    service = DayFactsService(store)
    pending = store.begin(17, date(2026, 10, 26))
    planner = SimpleNamespace(store=SimpleNamespace(session=lambda _id: None))
    notes = SimpleNamespace(pending_voice=lambda _id: False)
    router = create_day_facts_router(
        service=service, allowed_user_ids=frozenset({17}),
        planner=planner, notes=notes, export_dir=tmp_path,
    )
    handlers = {h.callback.__name__: h.callback for h in router.message.handlers}

    class MessageStub:
        chat = SimpleNamespace(type="private")
        from_user = SimpleNamespace(id=17)

        def __init__(self, text):
            self.text = text
            self.responses = []

        async def answer(self, value, **kwargs):
            self.responses.append((value, kwargs))

    for text, func_name in (
        ("📋 Текущие факты дня", "on_day"),
        ("🗓 История фактов", "on_history"),
    ):
        msg = MessageStub(text)
        await handlers[func_name](msg)
        assert msg.responses
        assert store.draft(17) == pending


@pytest.mark.asyncio
async def test_replace_command_corrects_text_without_redo(store, tmp_path):
    service = DayFactsService(store)
    pending = fill_record(store, 17, date(2026, 10, 26))
    planner = SimpleNamespace(store=SimpleNamespace(session=lambda _id: None))
    notes = SimpleNamespace(pending_voice=lambda _id: False)
    router = create_day_facts_router(
        service=service, allowed_user_ids=frozenset({17}),
        planner=planner, notes=notes, export_dir=tmp_path,
    )
    handlers = {h.callback.__name__: h.callback for h in router.message.handlers}

    class MessageStub:
        chat = SimpleNamespace(type="private")
        from_user = SimpleNamespace(id=17)
        text = (
            "/facts_replace facts | Друг общался раздражённо, "
            "был в своих мыслях. Мы разошлись по домам."
        )

        def __init__(self):
            self.responses = []

        async def answer(self, value, **kwargs):
            self.responses.append(value)

    msg = MessageStub()
    await handlers["on_replace"](msg)
    updated = store.draft(17)
    assert len(updated.facts) == 2
    assert updated.judgments == pending.judgments
    assert updated.stage == "review"
    assert any("Друг общался раздражённо" in value for value in msg.responses)



def test_existing_doubled_facts_draft_survives_upgrade_and_can_be_redone(tmp_path):
    """Regression for 26 Oct Telegram: user repeated the same voice twice."""
    path = tmp_path / "production-shaped.sqlite3"
    before = DayFactsStore(path)
    draft = before.begin(17, date(2026, 10, 26))
    draft = before.accept(
        17, draft.token, "Гуляли с анонимным другом, он был без настроения."
    )
    words = "Друг общался, раздраженно был в своих мыслях. Мы разошлись по домам."
    draft = before.accept(17, draft.token, words, raw_text=words)
    draft = before.accept(17, draft.token, words, raw_text=words)
    assert len(draft.facts) == 4
    before.close()

    upgraded = DayFactsStore(path)
    try:
        restored = upgraded.draft(17)
        assert restored.facts == draft.facts
        reset = upgraded.reset_section(17, restored.token, "facts")
        assert reset.facts == ()
        assert reset.event == restored.event
        assert upgraded.accept(17, restored.token, "queued old speech") is None
        corrected = upgraded.accept(
            17, reset.token,
            "Друг общался раздражённо, был в своих мыслях. Мы разошлись по домам.",
        )
        assert corrected.facts == (
            "Друг общался раздражённо, был в своих мыслях.",
            "Мы разошлись по домам.",
        )
        assert len(corrected.facts) == 2
    finally:
        upgraded.close()



def test_existing_live_sqlite_draft_is_migrated_without_loss(tmp_path):
    """A pre-update draft with four duplicate facts must survive schema upgrade."""
    from voqelis.day_facts.store import SCHEMA

    db_file = tmp_path / "day_facts.sqlite3"
    legacy_schema = SCHEMA.replace(
        "    updated_at TEXT NOT NULL,\n    edit_field TEXT,\n    edit_index INTEGER",
        "    updated_at TEXT NOT NULL",
    )
    assert legacy_schema != SCHEMA
    with sqlite3.connect(db_file) as connection:
        connection.executescript(legacy_schema)
        connection.execute(
            "INSERT INTO day_facts_drafts "
            "(user_id,day,stage,token,event,facts_json,judgments_json,updated_at) "
            "VALUES(?,?,?,?,?,?,?,?)",
            (
                71, "2026-10-26", "facts", "old-token",
                "Прогулка с другом",
                ('["Друг общался, раздраженно был в своих мыслях.",'
                 ' "Мы разошлись по домам.",'
                 ' "Друг общался, раздраженно был в своих мыслях.",'
                 ' "Мы разошлись по домам."]'),
                "[]", "2026-10-26T00:00:00Z",
            ),
        )
    reopened = DayFactsStore(db_file)
    try:
        draft = reopened.draft(71)
        assert draft is not None
        assert draft.edit_field is None
        assert len(draft.facts) == 4
        selected = reopened.select_edit(71, draft.token, "facts", 1)
        assert selected.edit_index == 1
        assert selected.edit_field == "facts"
        corrected = reopened.accept(
            71, selected.token, "Друг общался раздражённо, был в своих мыслях."
        )
        assert len(corrected.facts) == 4
        assert corrected.facts[0] == "Друг общался раздражённо, был в своих мыслях."
        assert corrected.facts[1:] == draft.facts[1:]
    finally:
        reopened.close()
    restarted = DayFactsStore(db_file)
    try:
        assert restarted.draft(71).facts[0] == (
            "Друг общался раздражённо, был в своих мыслях."
        )
    finally:
        restarted.close()


def test_single_item_edit_does_not_append_or_modify_other_items(store):
    original = fill_record(store, 71, date(2026, 10, 26))
    selection = store.select_edit(71, original.token, "facts", 1)
    assert selection.edit_field == "facts"
    assert selection.edit_index == 1
    assert selection.token != original.token
    assert store.accept(71, original.token, "очередь: устаревшее голосовое") is None
    replacement = store.accept(
        71, selection.token,
        "Друг общался раздражённо, был в своих мыслях.",
        raw_text="друг общался раздраженно был в своих мыслях",
    )
    assert replacement.stage == "review"
    assert replacement.edit_field is None
    assert replacement.facts == (
        "Друг общался раздражённо, был в своих мыслях.",
        "Мы разошлись по домам.",
    )
    assert replacement.judgments == original.judgments
    assert store.save(71, original.token) is None
    saved = store.save(71, replacement.token)
    assert saved.facts == replacement.facts


def test_one_edit_never_splits_into_more_items(store):
    d = fill_record(store, 71, date(2026, 10, 26))
    selected = store.select_edit(71, d.token, "judgments", 1)
    result = store.accept(
        71, selected.token, "Первая мысль. Вторая мысль."
    )
    assert len(result.judgments) == 2
    assert result.judgments[0] == "Первая мысль. Вторая мысль."
    assert result.judgments[1] == d.judgments[1]


def test_edit_cancellation_preserves_original_and_rejects_stale_voice(store):
    d = fill_record(store, 71, date(2026, 10, 26))
    pending = store.select_edit(71, d.token, "facts", 2)
    with pytest.raises(ValueError):
        store.next(71, pending.token)
    assert store.save(71, pending.token) is None
    with pytest.raises(ValueError):
        store.reset_section(71, pending.token, "facts")
    restored = store.cancel_edit(71, pending.token)
    assert restored.edit_field is None
    assert restored.facts == d.facts
    assert restored.stage == "review"
    assert store.accept(71, pending.token, "запоздавшее голосовое") is None
    assert store.save(71, restored.token) is not None


def test_select_event_and_judgment_and_invalid_indices(store):
    d = fill_record(store, 71, date(2026, 10, 26))
    e = store.select_edit(71, d.token, "event", 1)
    assert e.edit_field == "event"
    new = store.accept(71, e.token, "Другое событие.")
    assert new.event == "Другое событие."
    assert new.facts == d.facts
    j = store.select_edit(71, new.token, "judgments", 2)
    again = store.accept(71, j.token, "Я подумал, что я плохой друг.")
    assert again.judgments[1] == "Я подумал, что я плохой друг."
    with pytest.raises(ValueError):
        store.select_edit(71, again.token, "facts", 99)
    with pytest.raises(ValueError):
        store.select_edit(71, again.token, "event", 2)
    assert store.select_edit(72, again.token, "facts", 1) is None


def test_edit_buttons_are_intuitive_and_paginated_without_private_callback_text(store):
    from voqelis.day_facts.bot import (
        edit_items_markup,
        edit_sections_markup,
        facts_markup,
    )

    d = fill_record(store, 71, date(2026, 10, 26))
    buttons = [x.text for row in facts_markup(d).inline_keyboard for x in row]
    assert "✏️ Исправить" in buttons
    assert not any("/facts_" in label for label in buttons)
    fields = [
        x.text for row in edit_sections_markup(d).inline_keyboard for x in row
    ]
    assert "📝 Событие" in fields
    assert any("Факты" in name for name in fields)
    assert any("Суждения" in name for name in fields)
    many = store.replace_section(
        71, d.token, "facts",
        " ".join(f"Дополнительный факт номер {i}." for i in range(12)),
    )
    first = edit_items_markup(many, "facts", 0)
    all_buttons = [x for row in first.inline_keyboard for x in row]
    assert any(x.text == "➡️" for x in all_buttons)
    for button in all_buttons:
        assert "Дополнительный факт" not in button.callback_data
        assert len(button.callback_data) <= 64
    second = edit_items_markup(many, "facts", 1)
    assert any(
        x.text.startswith("9.") for row in second.inline_keyboard for x in row
    )
    with pytest.raises(ValueError):
        edit_items_markup(many, "facts", 42)
    selected = store.select_edit(71, many.token, "facts", 9)
    controls = [x.text for row in facts_markup(selected).inline_keyboard for x in row]
    assert controls == ["↩️ Отмена исправления"]


@pytest.mark.asyncio
async def test_edit_callback_button_replaces_only_selected_fact(store, tmp_path):
    from voqelis.day_facts.bot import create_day_facts_router

    d = fill_record(store, 71, date(2026, 10, 26))
    service = DayFactsService(store)
    planner = SimpleNamespace(store=SimpleNamespace(session=lambda _id: None))
    notes = SimpleNamespace(pending_voice=lambda _id: False)
    router = create_day_facts_router(
        service=service, allowed_user_ids=frozenset({71}),
        planner=planner, notes=notes, export_dir=tmp_path,
    )
    on_callback = router.callback_query.handlers[0].callback
    on_field = next(
        x.callback for x in router.message.handlers
        if x.callback.__name__ == "on_field_text"
    )

    class FakeMessage:
        from_user = SimpleNamespace(id=71)
        chat = SimpleNamespace(type="private")
        text = "Друг общался раздражённо, был в своих мыслях."

        def __init__(self):
            self.answers = []

        async def answer(self, body, **kw):
            self.answers.append((body, kw))

    class FakeCallback:
        from_user = SimpleNamespace(id=71)

        def __init__(self, data, message):
            self.message = message
            self.data = data
            self.acks = []

        async def answer(self, *args, **kwargs):
            self.acks.append((args, kwargs))

    msg = FakeMessage()
    first = FakeCallback(f"df:edit:{d.token}", msg)
    await on_callback(first)
    assert first.acks
    chooser = store.draft(71)
    assert chooser.edit_field == "select"
    assert chooser.token != d.token
    for data in (
        f"df:items:facts:0:{chooser.token}",
        f"df:pick:facts:1:{chooser.token}",
    ):
        cb = FakeCallback(data, msg)
        await on_callback(cb)
        assert cb.acks
    pending = store.draft(71)
    assert pending.edit_field == "facts"
    assert "Отправь новый текст или голосовое" in msg.answers[-1][0]
    await on_field(msg)
    final = store.draft(71)
    assert final.edit_field is None
    assert final.facts == (
        "Друг общался раздражённо, был в своих мыслях.",
        "Мы разошлись по домам.",
    )
    assert final.judgments == d.judgments
    assert len(final.facts) == 2


@pytest.mark.asyncio
async def test_voice_edit_uses_queued_token_and_returns_review(store, tmp_path, monkeypatch):
    service = DayFactsService(store)
    d = fill_record(store, 71, date(2026, 10, 26))
    pending = store.select_edit(71, d.token, "facts", 1)
    original = pending.token
    from voqelis import bot as bot_module

    monkeypatch.setattr(bot_module, "probe_duration_seconds", lambda _: 3)
    file = tmp_path / "edit.ogg"
    file.write_bytes(b"sample")
    bot = _FakeBot()
    job = AudioJob(
        chat_id=71, reply_to_message_id=1, user_id=71,
        file_path=file, day_facts_token=original,
    )
    with pytest.raises(asyncio.CancelledError):
        await run_worker(
            bot=bot, queue=_FakeQueue(job), transcriber=_FakeASR(),
            settings=SimpleNamespace(max_audio_seconds=90),
            day_facts=service, punctuation_processor=_FakePunctuation(),
        )
    now = store.draft(71)
    assert now.edit_field is None
    assert now.stage == "review"
    assert now.facts[0] == (
        "Друг общался раздражённо, был в своих мыслях. Мы разошлись по домам."
    )
    assert len(now.facts) == 2
    assert store.accept(71, original, "повтор") is None



@pytest.mark.asyncio
async def test_old_edit_command_without_arguments_opens_button_picker(store, tmp_path):
    """Telegram users should not see programming-like argument instructions."""
    from voqelis.day_facts.bot import create_day_facts_router

    d = fill_record(store, 71, date(2026, 10, 26))
    service = DayFactsService(store)
    planner = SimpleNamespace(store=SimpleNamespace(session=lambda _id: None))
    notes = SimpleNamespace(pending_voice=lambda _id: False)
    router = create_day_facts_router(
        service=service, allowed_user_ids=frozenset({71}),
        planner=planner, notes=notes, export_dir=tmp_path,
    )
    edit_handler = next(
        h.callback for h in router.message.handlers
        if h.callback.__name__ == "on_edit"
    )

    class FakeMessage:
        from_user = SimpleNamespace(id=71)
        chat = SimpleNamespace(type="private")
        text = "/facts_edit"

        def __init__(self):
            self.answers = []

        async def answer(self, body, **kwargs):
            self.answers.append((body, kwargs))

    message = FakeMessage()
    await edit_handler(message)
    assert len(message.answers) == 1
    prompt, kwargs = message.answers[0]
    assert "Выбери кнопкой" in prompt
    assert "/facts_edit facts" not in prompt
    controls = [
        btn.text
        for row in kwargs["reply_markup"].inline_keyboard
        for btn in row
    ]
    assert "📝 Событие" in controls
    assert any("Факты" in name for name in controls)
    chooser = store.draft(71)
    assert chooser.edit_field == "select"
    assert chooser.token != d.token
    assert chooser.facts == d.facts



def test_opening_edit_menu_blocks_append_until_item_is_selected(store):
    d = fill_record(store, 71, date(2026, 10, 26))
    chooser = store.open_edit_menu(71, d.token)
    assert chooser.edit_field == "select"
    assert chooser.edit_index is None
    assert store.accept(71, chooser.token, "ошибочно отправленный факт") is None
    assert store.accept(71, d.token, "старое голосовое из очереди") is None
    assert store.draft(71).facts == d.facts
    with pytest.raises(ValueError):
        store.next(71, chooser.token)
    assert store.save(71, chooser.token) is None
    picked = store.select_edit(71, chooser.token, "facts", 1)
    assert picked.edit_field == "facts"
    assert picked.token != chooser.token
    corrected = store.accept(
        71, picked.token, "Друг общался раздражённо, был в своих мыслях."
    )
    assert corrected.facts[0] == "Друг общался раздражённо, был в своих мыслях."
    assert corrected.facts[1] == d.facts[1]
    assert corrected.edit_field is None


@pytest.mark.asyncio
async def test_text_during_item_choice_is_rejected_not_appended(store, tmp_path):
    service = DayFactsService(store)
    d = fill_record(store, 71, date(2026, 10, 26))
    chooser = store.open_edit_menu(71, d.token)
    planner = SimpleNamespace(store=SimpleNamespace(session=lambda _id: None))
    notes = SimpleNamespace(pending_voice=lambda _id: False)
    router = create_day_facts_router(
        service=service, allowed_user_ids=frozenset({71}),
        planner=planner, notes=notes, export_dir=tmp_path,
    )
    on_field = next(
        h.callback for h in router.message.handlers
        if h.callback.__name__ == "on_field_text"
    )

    class FakeMessage:
        text = "Друг общался раздражённо, был в своих мыслях."
        chat = SimpleNamespace(type="private")
        from_user = SimpleNamespace(id=71)

        def __init__(self):
            self.responses = []

        async def answer(self, body, **kwargs):
            self.responses.append(body)

    msg = FakeMessage()
    await on_field(msg)
    assert "Сначала выбери раздел" in msg.responses[0]
    assert store.draft(71) == chooser
    assert store.draft(71).facts == d.facts


@pytest.mark.asyncio
async def test_voice_in_item_chooser_is_rejected_before_download_or_queue(store):
    from voqelis.bot import create_router

    service = DayFactsService(store)
    d = fill_record(store, 71, date(2026, 10, 26))
    chooser = store.open_edit_menu(71, d.token)

    class StubQueue:
        async def reserve(self, user_id):
            raise AssertionError("Voice must be blocked before queue reservation")

    router = create_router(
        settings=SimpleNamespace(
            allowed_user_ids=frozenset({71}),
            max_file_size_bytes=1000000,
            max_audio_seconds=60,
        ),
        queue=StubQueue(),
        day_facts=service,
    )
    on_audio = next(
        h.callback for h in router.message.handlers
        if h.callback.__name__ == "on_audio"
    )

    class FakeMessage:
        chat = SimpleNamespace(type="private")
        from_user = SimpleNamespace(id=71)
        voice = SimpleNamespace(file_id="fake", file_size=100, duration=3)
        audio = None
        document = None

        def __init__(self):
            self.replies = []

        async def reply(self, body):
            self.replies.append(body)

    msg = FakeMessage()
    await on_audio(msg, SimpleNamespace())
    assert len(msg.replies) == 1
    assert "Сначала выбери" in msg.replies[0]
    assert store.draft(71) == chooser


def test_cancel_edit_menu_restores_original_draft(store):
    d = fill_record(store, 71, date(2026, 10, 26))
    chooser = store.open_edit_menu(71, d.token)
    returned = store.cancel_edit(71, chooser.token)
    assert returned.edit_field is None
    assert returned.facts == d.facts
    assert returned.judgments == d.judgments
    assert store.accept(71, chooser.token, "поздняя диктовка") is None



def test_resuming_selector_shows_category_buttons_instead_of_dead_end(store):
    from voqelis.day_facts.bot import facts_markup

    draft = fill_record(store, 71, date(2026, 10, 26))
    selector = store.open_edit_menu(71, draft.token)
    buttons = [
        item.text
        for row in facts_markup(selector).inline_keyboard
        for item in row
    ]
    assert "📝 Событие" in buttons
    assert any("Факты" in item for item in buttons)
    assert "↩️ Назад" in buttons


@pytest.mark.asyncio
async def test_history_date_buttons_export_chosen_day_and_not_today(store, tmp_path):
    """Historical record PDF is selected by its callback date, never today's menu."""
    service = DayFactsService(store)
    service.today = lambda: date(2026, 10, 28)
    day = date(2026, 10, 26)
    draft = fill_record(store, 1, day)
    store.save(1, draft.token)

    planner = SimpleNamespace(store=SimpleNamespace(session=lambda _id: None))
    notes = SimpleNamespace(pending_voice=lambda _id: False)
    router = create_day_facts_router(
        service=service, allowed_user_ids=frozenset({1, 2}),
        planner=planner, notes=notes, export_dir=tmp_path,
    )

    class HistoryMessage:
        def __init__(self, user_id, text=""):
            self.text = text
            self.from_user = SimpleNamespace(id=user_id)
            self.chat = SimpleNamespace(type="private")
            self.responses = []
            self.documents = []

        async def answer(self, value, **kwargs):
            self.responses.append((value, kwargs))

        async def answer_document(self, attachment, **kwargs):
            self.documents.append((
                attachment.filename, Path(attachment.path).read_bytes(),
            ))

    class Click:
        def __init__(self, user_id, data, message):
            self.from_user = SimpleNamespace(id=user_id)
            self.data = data
            self.message = message
            self.acks = []

        async def answer(self, text=None, **kwargs):
            self.acks.append((text, kwargs))

    handlers = {h.callback.__name__: h.callback for h in router.message.handlers}
    handle_callback = router.callback_query.handlers[0].callback

    history = HistoryMessage(1, "🗓 История фактов")
    await handlers["on_history"](history)
    assert "26.10.2026" in history.responses[-1][0]
    picker = history.responses[-1][1]["reply_markup"]
    assert picker.inline_keyboard[0][0].callback_data == "df:day:2026-10-26"

    opened = Click(1, "df:day:2026-10-26", history)
    await handle_callback(opened)
    assert opened.acks
    assert "Друг общался" in history.responses[-1][0]
    markup = history.responses[-1][1]["reply_markup"]
    assert [
        button.callback_data
        for row in markup.inline_keyboard for button in row
    ] == [
        "df:export:word:2026-10-26",
        "df:export:pdf:2026-10-26",
        "df:export:word_mobile:2026-10-26",
        "df:export:mobile:2026-10-26",
    ]

    pdf_click = Click(1, "df:export:pdf:2026-10-26", history)
    await handle_callback(pdf_click)
    assert len(history.documents) == 1
    assert history.documents[0][0] == "fakty-dnya-2026-10-26-pdf.pdf"
    assert history.documents[0][1].startswith(b"%PDF")

    direct = HistoryMessage(1, "/facts_day 2026-10-26")
    await handlers["on_day"](direct)
    assert direct.responses[-1][1]["reply_markup"] is not None

    today = HistoryMessage(1, "/facts_day 2026-10-28")
    await handlers["on_day"](today)
    assert today.responses[-1][1]["reply_markup"] is None


@pytest.mark.asyncio
async def test_history_and_forged_export_cannot_read_another_users_facts(store, tmp_path):
    """Button payloads are not authorization: sender ID scopes every lookup."""
    service = DayFactsService(store)
    day = date(2026, 10, 26)
    draft = fill_record(store, 1, day)
    store.save(1, draft.token)

    planner = SimpleNamespace(store=SimpleNamespace(session=lambda _id: None))
    notes = SimpleNamespace(pending_voice=lambda _id: False)
    router = create_day_facts_router(
        service=service, allowed_user_ids=frozenset({1, 2}),
        planner=planner, notes=notes, export_dir=tmp_path,
    )
    handlers = {h.callback.__name__: h.callback for h in router.message.handlers}
    callback_handler = router.callback_query.handlers[0].callback

    class GuestMessage:
        text = "/facts_history"
        chat = SimpleNamespace(type="private")
        from_user = SimpleNamespace(id=2)

        def __init__(self):
            self.responses = []
            self.documents = []

        async def answer(self, value, **kwargs):
            self.responses.append((value, kwargs))

        async def answer_document(self, attachment, **kwargs):
            self.documents.append(attachment.filename)

    class GuestClick:
        message = GuestMessage()
        from_user = SimpleNamespace(id=2)

        def __init__(self, data):
            self.data = data
            self.answers = []

        async def answer(self, value=None, **kwargs):
            self.answers.append((value, kwargs))

    guest = GuestMessage()
    await handlers["on_history"](guest)
    assert "Записей пока нет" in guest.responses[-1][0]
    assert guest.responses[-1][1]["reply_markup"] is None

    forged_day = GuestClick("df:day:2026-10-26")
    await callback_handler(forged_day)
    assert "сохранённых ситуаций нет" in forged_day.message.responses[-1][0]
    assert forged_day.message.responses[-1][1]["reply_markup"] is None

    forged_export = GuestClick("df:export:word_mobile:2026-10-26")
    await callback_handler(forged_export)
    assert "Нет сохранённых ситуаций" in forged_export.message.responses[-1][0]
    assert forged_export.message.documents == []

    bad_date = GuestClick("df:day:bad-date")
    await callback_handler(bad_date)
    assert bad_date.answers[-1][1].get("show_alert") is True

